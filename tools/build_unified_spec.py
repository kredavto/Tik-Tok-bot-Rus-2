from __future__ import annotations

import re
import textwrap
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Flowable,
    Frame,
    PageBreak,
    PageTemplate,
    Paragraph,
    Preformatted,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT_DIR = DOCS / "final"

PROJECT_TITLE = "Tik_Tok_Loader"
DOC_TITLE_RU = "Единая техническая спецификация"
DOC_TITLE_EN = "Unified Technical Specification"
VERSION = "0.1.0"
REPO = "kredavto/Tik-Tok-bot-Rus-2"


@dataclass(frozen=True)
class SourceDoc:
    path: str
    group: str
    appendix: bool = False


SOURCE_DOCS: list[SourceDoc] = [
    SourceDoc("compliance.md", "Общие положения"),
    SourceDoc("architecture-summary.md", "Архитектура"),
    SourceDoc("component-map.md", "Архитектура"),
    SourceDoc("sequence-flows.md", "Архитектура"),
    SourceDoc("implementation-roadmap.md", "Управление реализацией"),
    SourceDoc("non-functional-requirements.md", "Нефункциональные требования"),
    SourceDoc("data-model.md", "Данные"),
    SourceDoc("users-entity.md", "Данные"),
    SourceDoc("tiktok-accounts-entity.md", "Данные"),
    SourceDoc("subscriptions-entity.md", "Данные"),
    SourceDoc("payments-entity.md", "Данные"),
    SourceDoc("upload-jobs-entity.md", "Данные"),
    SourceDoc("daily-usage-entity.md", "Данные"),
    SourceDoc("webhook-events-entity.md", "Данные"),
    SourceDoc("admin-actions-entity.md", "Данные"),
    SourceDoc("system-settings-entity.md", "Данные"),
    SourceDoc("user-guide.md", "Пользовательские сценарии"),
    SourceDoc("video-lifecycle.md", "Пользовательские сценарии"),
    SourceDoc("tiktok-developer-configuration.md", "Интеграции"),
    SourceDoc("robokassa.md", "Интеграции"),
    SourceDoc("api.md", "API"),
    SourceDoc("rest-api-standards.md", "API"),
    SourceDoc("api-versioning-compatibility.md", "API"),
    SourceDoc("public-rest-api.md", "API"),
    SourceDoc("admin-rest-api.md", "API"),
    SourceDoc("openapi-contracts.md", "API"),
    SourceDoc("error-handling.md", "API"),
    SourceDoc("security.md", "Безопасность"),
    SourceDoc("security-logging-audit.md", "Безопасность"),
    SourceDoc("confidential-data-policy.md", "Безопасность"),
    SourceDoc("environment-configuration-secrets.md", "Безопасность"),
    SourceDoc("configuration.md", "Конфигурация"),
    SourceDoc("configuration-management.md", "Конфигурация"),
    SourceDoc("deployment.md", "Эксплуатация"),
    SourceDoc("containerization.md", "Эксплуатация"),
    SourceDoc("production-launch.md", "Эксплуатация"),
    SourceDoc("operations-runbook.md", "Эксплуатация"),
    SourceDoc("sop-checklists.md", "Эксплуатация"),
    SourceDoc("maintenance.md", "Эксплуатация"),
    SourceDoc("post-launch-maintenance.md", "Эксплуатация"),
    SourceDoc("incident-response.md", "Эксплуатация"),
    SourceDoc("incident-management.md", "Эксплуатация"),
    SourceDoc("backup-restore-policy.md", "Надежность"),
    SourceDoc("service-continuity-plan.md", "Надежность"),
    SourceDoc("data-retention.md", "Надежность"),
    SourceDoc("file-storage-policy.md", "Надежность"),
    SourceDoc("queue-retry-policy.md", "Надежность"),
    SourceDoc("performance.md", "Производительность"),
    SourceDoc("capacity-management.md", "Производительность"),
    SourceDoc("metrics-and-kpi.md", "Производительность"),
    SourceDoc("observability-diagnostics.md", "Наблюдаемость"),
    SourceDoc("development.md", "Разработка"),
    SourceDoc("github-workflow.md", "Разработка"),
    SourceDoc("feature-development-plan.md", "Разработка"),
    SourceDoc("technical-debt-management.md", "Разработка"),
    SourceDoc("dependencies-and-integrations.md", "Зависимости"),
    SourceDoc("infrastructure-dependency-management.md", "Зависимости"),
    SourceDoc("license-third-party-management.md", "Зависимости"),
    SourceDoc("migration-compatibility-plan.md", "Релизы"),
    SourceDoc("release.md", "Релизы"),
    SourceDoc("change-acceptance-policy.md", "Релизы"),
    SourceDoc("qa-acceptance-scenarios.md", "Качество"),
    SourceDoc("acceptance-checklist.md", "Качество"),
    SourceDoc("requirements-traceability.md", "Приложения", appendix=True),
    SourceDoc("glossary-naming.md", "Приложения", appendix=True),
    SourceDoc("specification-index.md", "Приложения", appendix=True),
    SourceDoc("risk-management.md", "Приложения", appendix=True),
]


ABBREVIATIONS: list[tuple[str, str]] = [
    ("API", "Application Programming Interface, программный интерфейс приложения"),
    ("CI/CD", "Continuous Integration / Continuous Delivery"),
    ("CSRF", "Cross-Site Request Forgery"),
    ("FSM", "Finite State Machine, конечный автомат сценариев Telegram-бота"),
    ("JSON", "JavaScript Object Notation"),
    ("KPI", "Key Performance Indicator"),
    ("NFR", "Non-Functional Requirements, нефункциональные требования"),
    ("OAuth 2.0", "Протокол авторизации для подключения аккаунта TikTok"),
    ("RBAC", "Role-Based Access Control, ролевая модель доступа"),
    ("REST", "Representational State Transfer"),
    ("SOP", "Standard Operating Procedure, стандартная операционная процедура"),
    ("TLS/SSL", "Криптографическая защита транспортного соединения"),
    ("UTC", "Coordinated Universal Time"),
]


def slugify(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[`*_{}\[\]()+.!?,:;/\\]", "", value)
    value = re.sub(r"\s+", "-", value)
    return value


def strip_md_links(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    return text


def read_title(path: Path) -> str:
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            return strip_md_links(line[2:].strip())
    return path.stem.replace("-", " ").title()


def normalize_text(text: str) -> str:
    text = text.replace("→", "->")
    text = text.replace("—", "-")
    text = text.replace("–", "-")
    text = text.replace("‑", "-")
    return text


def table_row(line: str) -> list[str] | None:
    if not line.strip().startswith("|") or not line.strip().endswith("|"):
        return None
    if re.fullmatch(r"\s*\|[\s:\-|\u2014]+\|\s*", line):
        return []
    return [strip_md_links(cell.strip()) for cell in line.strip().strip("|").split("|")]


def add_field(paragraph, instruction: str, placeholder: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)

    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_separate)

    paragraph.add_run(placeholder)

    run_end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, column_count: int) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = column_widths(column_count)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")


def column_widths(column_count: int) -> list[int]:
    if column_count <= 1:
        return [9360]
    if column_count == 2:
        return [2600, 6760]
    if column_count == 3:
        return [2200, 2960, 4200]
    if column_count == 4:
        return [1500, 2400, 2760, 2700]
    if column_count == 5:
        return [1200, 2200, 2200, 1800, 1960]
    width = 9360 // column_count
    return [width] * column_count


def pdf_column_widths(column_count: int) -> list[float]:
    widths_dxa = column_widths(column_count)
    return [6.5 * inch * (width / 9360) for width in widths_dxa]


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = f"{PROJECT_TITLE} - {DOC_TITLE_RU}"
        header.style = doc.styles["Header"]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Страница ")
        add_field(footer, "PAGE", "1")


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(PROJECT_TITLE)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string("1F4D78")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"{DOC_TITLE_RU}\n{DOC_TITLE_EN}")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string("2E74B5")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Версия: {VERSION}\n"
        f"Репозиторий: {REPO}\n"
        f"Дата сборки: {datetime.now(UTC).strftime('%Y-%m-%d')}\n"
        "Статус: проектная спецификация для реализации"
    )

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "Публикация TikTok в проекте проектируется только через официальный "
        "TikTok Content Posting API и OAuth 2.0. Неофициальные API, автоматизация "
        "интерфейса и методы обхода ограничений не входят в допустимую архитектуру."
    )
    run.italic = True
    run.font.size = Pt(10)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Оглавление")
    p = doc.add_paragraph()
    add_field(p, r'TOC \o "1-3" \h \z \u', "Оглавление обновляется в Word/LibreOffice.")
    doc.add_page_break()


def add_abbreviations(doc: Document) -> None:
    doc.add_heading("Список сокращений и терминов", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Сокращение / термин"
    table.rows[0].cells[1].text = "Описание"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "E8EEF5")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for key, value in ABBREVIATIONS:
        row = table.add_row().cells
        row[0].text = key
        row[1].text = value
    set_table_geometry(table, 2)
    doc.add_page_break()


def emit_md_front_matter() -> list[str]:
    generated = datetime.now(UTC).strftime("%Y-%m-%d")
    lines = [
        f"# {PROJECT_TITLE}. {DOC_TITLE_RU}",
        "",
        f"**{DOC_TITLE_EN}**",
        "",
        f"- **Версия:** {VERSION}",
        f"- **Репозиторий:** `{REPO}`",
        f"- **Дата сборки:** {generated}",
        "- **Статус:** проектная спецификация для реализации",
        "",
        "> Публикация TikTok в проекте проектируется только через официальный TikTok Content Posting API и OAuth 2.0. Неофициальные API, автоматизация интерфейса и методы обхода ограничений не входят в допустимую архитектуру.",
        "",
        "## Оглавление",
        "",
    ]
    for number, source in enumerate(SOURCE_DOCS, start=1):
        title = read_title(DOCS / source.path)
        lines.append(f"- [{number}. {title}](#{slugify(f'{number} {title}')})")
    lines.extend(["", "## Список сокращений и терминов", ""])
    lines.append("| Сокращение / термин | Описание |")
    lines.append("| --- | --- |")
    for key, value in ABBREVIATIONS:
        lines.append(f"| {key} | {value} |")
    lines.append("")
    return lines


def convert_markdown(source_docs: Iterable[SourceDoc]) -> str:
    lines = emit_md_front_matter()
    for idx, source in enumerate(source_docs, start=1):
        path = DOCS / source.path
        content = normalize_text(path.read_text(encoding="utf-8")).splitlines()
        h2_count = 0
        h3_count = 0
        h4_count = 0
        in_code = False
        first_h1_done = False
        title = read_title(path)
        lines.extend(["", f"# {idx}. {title}", ""])
        lines.append(f"Группа спецификации: {source.group}")
        lines.append("")
        for raw in content:
            line = raw.rstrip()
            if line.startswith("```"):
                in_code = not in_code
                lines.append(line)
                continue
            if in_code:
                lines.append(line)
                continue
            if line.startswith("# "):
                if first_h1_done:
                    h2_count += 1
                    h3_count = 0
                    lines.append(f"## {idx}.{h2_count}. {strip_md_links(line[2:].strip())}")
                    continue
                first_h1_done = True
                continue
            if line.startswith("## "):
                h2_count += 1
                h3_count = 0
                h4_count = 0
                lines.append(f"## {idx}.{h2_count}. {strip_md_links(line[3:].strip())}")
                continue
            if line.startswith("### "):
                h3_count += 1
                h4_count = 0
                lines.append(f"### {idx}.{h2_count}.{h3_count}. {strip_md_links(line[4:].strip())}")
                continue
            if line.startswith("#### "):
                h4_count += 1
                lines.append(
                    f"#### {idx}.{h2_count}.{h3_count}.{h4_count}. {strip_md_links(line[5:].strip())}"
                )
                continue
            lines.append(line)
    lines.append("")
    return "\n".join(lines)


def add_markdown_to_docx(doc: Document, markdown: str) -> None:
    code_lines: list[str] = []
    table_lines: list[list[str]] = []
    in_code = False

    def flush_code() -> None:
        nonlocal code_lines
        if not code_lines:
            return
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        code_lines = []

    def flush_table() -> None:
        nonlocal table_lines
        if not table_lines:
            return
        max_cols = max(len(row) for row in table_lines)
        table = doc.add_table(rows=1, cols=max_cols)
        table.style = "Table Grid"
        for idx, cell in enumerate(table.rows[0].cells):
            cell.text = table_lines[0][idx] if idx < len(table_lines[0]) else ""
            set_cell_shading(cell, "E8EEF5")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        for row_values in table_lines[1:]:
            row = table.add_row().cells
            for idx, cell in enumerate(row):
                cell.text = row_values[idx] if idx < len(row_values) else ""
        set_table_geometry(table, max_cols)
        doc.add_paragraph()
        table_lines = []

    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            flush_table()
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        row = table_row(line)
        if row is not None:
            if row:
                table_lines.append(row)
            continue
        flush_table()
        if not line.strip():
            continue
        if line.startswith("# "):
            title = strip_md_links(line[2:].strip())
            if title != f"{PROJECT_TITLE}. {DOC_TITLE_RU}":
                doc.add_section(WD_SECTION_START.NEW_PAGE)
            doc.add_heading(title, level=1)
            continue
        if line.startswith("## "):
            doc.add_heading(strip_md_links(line[3:].strip()), level=2)
            continue
        if line.startswith("### "):
            doc.add_heading(strip_md_links(line[4:].strip()), level=3)
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph()
            run = p.add_run(strip_md_links(line[5:].strip()))
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("1F4D78")
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(strip_md_links(line[2:].strip()))
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.add_run(strip_md_links(numbered.group(2)))
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(strip_md_links(line[2:].strip()))
            run.italic = True
            run.font.color.rgb = RGBColor.from_string("555555")
            continue
        p = doc.add_paragraph()
        p.add_run(strip_md_links(line))
    flush_table()
    flush_code()


def build_docx(markdown: str, output_path: Path) -> None:
    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_toc(doc)
    add_abbreviations(doc)
    # Skip title/front-matter from markdown because it is represented by cover, TOC, abbreviations.
    content_start = markdown.find("\n# 1. ")
    content_md = markdown[content_start + 1 :] if content_start != -1 else markdown
    add_markdown_to_docx(doc, content_md)
    add_header_footer(doc)
    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    doc.save(output_path)


class PageNumCanvasDoc(BaseDocTemplate):
    def __init__(self, filename: str, **kwargs):
        super().__init__(filename, **kwargs)
        frame = Frame(
            self.leftMargin,
            self.bottomMargin,
            self.width,
            self.height,
            id="normal",
        )
        template = PageTemplate(id="spec", frames=[frame], onPage=self.draw_page)
        self.addPageTemplates([template])

    def draw_page(self, canvas, doc) -> None:  # noqa: ANN001
        canvas.saveState()
        canvas.setFont("ArialUnicode", 8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawRightString(
            LETTER[0] - doc.rightMargin,
            LETTER[1] - 0.55 * inch,
            f"{PROJECT_TITLE} - {DOC_TITLE_RU}",
        )
        canvas.drawCentredString(LETTER[0] / 2, 0.55 * inch, f"Страница {doc.page}")
        canvas.restoreState()


def register_pdf_fonts() -> None:
    candidates = [
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            pdfmetrics.registerFont(TTFont("ArialUnicode", str(candidate)))
            return
    raise RuntimeError("No suitable TrueType font found for PDF generation")


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "SpecTitle",
            parent=base["Title"],
            fontName="ArialUnicode",
            fontSize=28,
            leading=34,
            textColor=colors.HexColor("#1F4D78"),
            alignment=TA_CENTER,
            spaceAfter=18,
        ),
        "subtitle": ParagraphStyle(
            "SpecSubtitle",
            parent=base["Normal"],
            fontName="ArialUnicode",
            fontSize=15,
            leading=20,
            textColor=colors.HexColor("#2E74B5"),
            alignment=TA_CENTER,
            spaceAfter=20,
        ),
        "body": ParagraphStyle(
            "SpecBody",
            parent=base["Normal"],
            fontName="ArialUnicode",
            fontSize=9.3,
            leading=12.3,
            alignment=TA_LEFT,
            spaceAfter=6,
        ),
        "muted": ParagraphStyle(
            "SpecMuted",
            parent=base["Normal"],
            fontName="ArialUnicode",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#555555"),
            alignment=TA_LEFT,
            leftIndent=12,
            spaceAfter=8,
        ),
        "h1": ParagraphStyle(
            "SpecH1",
            parent=base["Heading1"],
            fontName="ArialUnicode",
            fontSize=15,
            leading=19,
            textColor=colors.HexColor("#2E74B5"),
            spaceBefore=18,
            spaceAfter=9,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "SpecH2",
            parent=base["Heading2"],
            fontName="ArialUnicode",
            fontSize=12.5,
            leading=16,
            textColor=colors.HexColor("#2E74B5"),
            spaceBefore=12,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "h3": ParagraphStyle(
            "SpecH3",
            parent=base["Heading3"],
            fontName="ArialUnicode",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#1F4D78"),
            spaceBefore=9,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "code": ParagraphStyle(
            "SpecCode",
            parent=base["Code"],
            fontName="Courier",
            fontSize=7,
            leading=8.5,
            leftIndent=8,
            rightIndent=8,
            spaceBefore=4,
            spaceAfter=6,
        ),
        "toc": ParagraphStyle(
            "SpecToc",
            parent=base["Normal"],
            fontName="ArialUnicode",
            fontSize=8.4,
            leading=10.5,
            spaceAfter=2,
        ),
    }


def para(text: str, style: ParagraphStyle) -> Paragraph:
    text = strip_md_links(text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"`([^`]+)`", r"<font name='Courier'>\1</font>", text)
    return Paragraph(escape(text, {"<b>": "<b>", "</b>": "</b>", "<font name='Courier'>": "<font name='Courier'>", "</font>": "</font>"}), style)


def clean_pdf_text(text: str) -> str:
    return strip_md_links(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(markdown: str, output_path: Path) -> None:
    register_pdf_fonts()
    styles = pdf_styles()
    doc = PageNumCanvasDoc(
        str(output_path),
        pagesize=LETTER,
        leftMargin=1 * inch,
        rightMargin=1 * inch,
        topMargin=0.82 * inch,
        bottomMargin=0.82 * inch,
    )
    story: list[Flowable] = []

    story.extend(
        [
            Spacer(1, 2.0 * inch),
            Paragraph(PROJECT_TITLE, styles["title"]),
            Paragraph(f"{DOC_TITLE_RU}<br/>{DOC_TITLE_EN}", styles["subtitle"]),
            Paragraph(
                f"Версия: {VERSION}<br/>Репозиторий: {REPO}<br/>"
                f"Дата сборки: {datetime.now(UTC).strftime('%Y-%m-%d')}<br/>"
                "Статус: проектная спецификация для реализации",
                styles["subtitle"],
            ),
            Spacer(1, 0.3 * inch),
            Paragraph(
                "Публикация TikTok в проекте проектируется только через официальный "
                "TikTok Content Posting API и OAuth 2.0. Неофициальные API, автоматизация "
                "интерфейса и методы обхода ограничений не входят в допустимую архитектуру.",
                styles["muted"],
            ),
            PageBreak(),
            Paragraph("Оглавление", styles["h1"]),
        ]
    )
    for number, source in enumerate(SOURCE_DOCS, start=1):
        title = read_title(DOCS / source.path)
        story.append(Paragraph(f"{number}. {clean_pdf_text(title)}", styles["toc"]))
    story.extend([PageBreak(), Paragraph("Список сокращений и терминов", styles["h1"])])

    abbrev_data = [["Сокращение / термин", "Описание"], *ABBREVIATIONS]
    story.append(
        Table(
            [[Paragraph(clean_pdf_text(str(cell)), styles["body"]) for cell in row] for row in abbrev_data],
            colWidths=pdf_column_widths(2),
            repeatRows=1,
            style=TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8D3DF")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            ),
        )
    )
    story.append(PageBreak())

    content_start = markdown.find("\n# 1. ")
    content_md = markdown[content_start + 1 :] if content_start != -1 else markdown
    in_code = False
    code_lines: list[str] = []
    table_lines: list[list[str]] = []

    def flush_code() -> None:
        nonlocal code_lines
        if not code_lines:
            return
        wrapped: list[str] = []
        for line in code_lines:
            wrapped.extend(textwrap.wrap(line, width=105, replace_whitespace=False) or [""])
        story.append(Preformatted("\n".join(wrapped), styles["code"]))
        code_lines = []

    def flush_table() -> None:
        nonlocal table_lines
        if not table_lines:
            return
        max_cols = max(len(row) for row in table_lines)
        data = []
        for row in table_lines:
            padded = row + [""] * (max_cols - len(row))
            data.append([Paragraph(clean_pdf_text(cell), styles["body"]) for cell in padded])
        story.append(
            Table(
                data,
                colWidths=pdf_column_widths(max_cols),
                repeatRows=1,
                style=TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8D3DF")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                ),
            )
        )
        story.append(Spacer(1, 8))
        table_lines = []

    first_h1 = True
    for raw in content_md.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            flush_table()
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        row = table_row(line)
        if row is not None:
            if row:
                table_lines.append(row)
            continue
        flush_table()
        if not line.strip():
            continue
        if line.startswith("# "):
            if first_h1:
                first_h1 = False
            else:
                story.append(PageBreak())
            story.append(Paragraph(clean_pdf_text(line[2:].strip()), styles["h1"]))
        elif line.startswith("## "):
            story.append(Paragraph(clean_pdf_text(line[3:].strip()), styles["h2"]))
        elif line.startswith("### "):
            story.append(Paragraph(clean_pdf_text(line[4:].strip()), styles["h3"]))
        elif line.startswith("#### "):
            story.append(Paragraph(f"<b>{clean_pdf_text(line[5:].strip())}</b>", styles["body"]))
        elif line.startswith("- "):
            story.append(Paragraph(f"• {clean_pdf_text(line[2:].strip())}", styles["body"]))
        elif re.match(r"^\d+\.\s+", line):
            story.append(Paragraph(clean_pdf_text(line), styles["body"]))
        elif line.startswith("> "):
            story.append(Paragraph(clean_pdf_text(line[2:].strip()), styles["muted"]))
        else:
            story.append(Paragraph(clean_pdf_text(line), styles["body"]))
    flush_table()
    flush_code()
    doc.build(story)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    markdown = convert_markdown(SOURCE_DOCS)
    md_path = OUTPUT_DIR / "Tik_Tok_Loader_Unified_Specification.md"
    docx_path = OUTPUT_DIR / "Tik_Tok_Loader_Unified_Specification.docx"
    pdf_path = OUTPUT_DIR / "Tik_Tok_Loader_Unified_Specification.pdf"
    md_path.write_text(markdown, encoding="utf-8")
    build_docx(markdown, docx_path)
    build_pdf(markdown, pdf_path)
    print(md_path)
    print(docx_path)
    print(pdf_path)


if __name__ == "__main__":
    main()
